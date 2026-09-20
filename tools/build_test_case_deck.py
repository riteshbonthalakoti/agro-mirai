"""Builds TEST_CASES.pptx and TEST_CASES.md from the real results in test_results.json.

    python tools/build_test_case_deck.py <test_results.json> <output folder>

Same table layout as the team's reference deck: Slno | Module | Test case |
Expected result | Status | Work Done, one table per test level (unit,
integration, system) for the backend and for the models.
"""
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

results = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
out = Path(sys.argv[2])
out.mkdir(parents=True, exist_ok=True)

GREEN = RGBColor(0x3F, 0x6B, 0x3A)
LIGHT = RGBColor(0xE7, 0xEF, 0xE2)
INK = RGBColor(0x22, 0x22, 0x22)
FONT = "Times New Roman"
ROWS_PER_SLIDE = 4
LEVEL_NAME = {"Unit": "Unit Testing", "Integration": "Integration Testing", "System": "System Testing"}
FIRST_COL = {"Unit": "Module", "Integration": "Integrated Modules", "System": "System Function"}
NUMBER = {("Backend", "Unit"): "13.1", ("Backend", "Integration"): "13.2", ("Backend", "System"): "13.3",
          ("Models", "Unit"): "13.4", ("Models", "Integration"): "13.5", ("Models", "System"): "13.6"}


def status_text(r):
    return "Completed" if r["status"] == "PASS" else "Failed"


def work_done(r):
    text = r["observed"].replace("\n", " ")
    if len(text) > 150:
        text = text[:147].rsplit(" ", 1)[0] + "..."
    return "Tested live: " + text


prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]


def add_text(slide, text, left, top, width, height, size, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    for run in p.runs:
        run.font.size, run.font.bold, run.font.name, run.font.color.rgb = Pt(size), bold, FONT, color
    return box


def footer(slide, n):
    add_text(slide, "AGRO MIRAI, Dept of AIML", Inches(0.4), Inches(7.05), Inches(6), Inches(0.35), 11, color=RGBColor(0x66, 0x66, 0x66))
    add_text(slide, str(n), Inches(12.2), Inches(7.05), Inches(0.8), Inches(0.35), 11, color=RGBColor(0x66, 0x66, 0x66), align=PP_ALIGN.RIGHT)


def set_cell(cell, text, size=12, bold=False, color=INK, fill=None, align=None):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    for run in p.runs:
        run.font.size, run.font.bold, run.font.name, run.font.color.rgb = Pt(size), bold, FONT, color
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    cell.margin_left = cell.margin_right = Inches(0.06)
    cell.margin_top = cell.margin_bottom = Inches(0.04)


def table_slide(title, first_col, chunk, start_no):
    s = prs.slides.add_slide(blank)
    add_text(s, title, Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.7), 26, bold=True, color=GREEN)
    widths = [0.6, 2.0, 2.9, 3.0, 1.1, 3.2]
    heads = ["Slno", first_col, "Test case", "Expected result", "Status", "Work Done"]
    shape = s.shapes.add_table(len(chunk) + 1, 6, Inches(0.4), Inches(1.05), Inches(sum(widths)), Inches(0.5 + 1.25 * len(chunk)))
    tbl = shape.table
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(0.45)
    for ri in range(1, len(chunk) + 1):
        tbl.rows[ri].height = Inches(1.15)
    for i, h in enumerate(heads):
        set_cell(tbl.cell(0, i), h, 13, True, RGBColor(255, 255, 255), GREEN)
    for ri, r in enumerate(chunk, start=1):
        fill = LIGHT if ri % 2 == 0 else RGBColor(255, 255, 255)
        vals = [f"{start_no + ri - 1}.", r["module"], r["test"], r["expected"], status_text(r), work_done(r)]
        for ci, v in enumerate(vals):
            set_cell(tbl.cell(ri, ci), v, 12, ci == 4, INK, fill)
    footer(s, len(prs.slides))


# ---- title slide
s = prs.slides.add_slide(blank)
add_text(s, "AGRO MIRAI", Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.0), 44, True, GREEN, PP_ALIGN.CENTER)
add_text(s, "13. Testing: Test Cases for the Backend and the Models", Inches(0.8), Inches(3.1), Inches(11.7), Inches(0.9), 28, True, INK, PP_ALIGN.CENTER)
add_text(s, "AI-driven smart agriculture advisory  |  BITM, Dept of AIML  |  Every test case below was run live on the demo backend",
         Inches(0.8), Inches(4.2), Inches(11.7), Inches(0.9), 16, False, RGBColor(0x55, 0x55, 0x55), PP_ALIGN.CENTER)
footer(s, 1)

# ---- summary slide
s = prs.slides.add_slide(blank)
add_text(s, "13. Testing: Summary of Results", Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.7), 26, True, GREEN)
areas = ["Backend", "Models"]
levels = ["Unit", "Integration", "System"]
shape = s.shapes.add_table(len(areas) + 2, 5, Inches(0.4), Inches(1.3), Inches(12.5), Inches(2.6))
t = shape.table
for i, w in enumerate([3.0, 2.4, 2.4, 2.4, 2.3]):
    t.columns[i].width = Inches(w)
for i, h in enumerate(["Area", "Unit tests", "Integration tests", "System tests", "Total passed"]):
    set_cell(t.cell(0, i), h, 14, True, RGBColor(255, 255, 255), GREEN)
grand_p = grand_n = 0
for ri, a in enumerate(areas, start=1):
    tp = tn = 0
    set_cell(t.cell(ri, 0), a, 14, True)
    for ci, lv in enumerate(levels, start=1):
        rs = [r for r in results if r["area"] == a and r["level"] == lv]
        p = sum(r["status"] == "PASS" for r in rs)
        tp += p; tn += len(rs)
        set_cell(t.cell(ri, ci), f"{p} / {len(rs)}", 14)
    set_cell(t.cell(ri, 4), f"{tp} / {tn}", 14, True)
    grand_p += tp; grand_n += tn
set_cell(t.cell(3, 0), "All", 14, True, fill=LIGHT)
for ci, lv in enumerate(levels, start=1):
    rs = [r for r in results if r["level"] == lv]
    set_cell(t.cell(3, ci), f"{sum(r['status'] == 'PASS' for r in rs)} / {len(rs)}", 14, fill=LIGHT)
set_cell(t.cell(3, 4), f"{grand_p} / {grand_n}", 14, True, fill=LIGHT)
add_text(s, "Unit testing checks one module alone, integration testing checks modules working together "
            "(for example live data adapters with the database, or the CNN with the API), and system testing "
            "checks the whole backend and all three models end to end.", Inches(0.4), Inches(4.3), Inches(12.5), Inches(1.4), 16)
add_text(s, "Test method: a script (test_cases.py) calls the running backend and the models with real requests and real "
            "live data (Open-Meteo, SoilGrids, Google Earth Engine) and records what actually happened.",
         Inches(0.4), Inches(5.6), Inches(12.5), Inches(1.2), 14, color=RGBColor(0x55, 0x55, 0x55))
footer(s, 2)

# ---- table slides + markdown
md = ["# AGRO MIRAI: Test Cases (Backend and Models)", "",
      f"Every case was executed live on the demo backend: **{grand_p} of {grand_n} passed**. "
      "Columns follow the reference deck: Slno, Module, Test case, Expected result, Status, Work Done.", ""]
for a in areas:
    md += [f"## {a}", ""]
    for lv in levels:
        rs = [r for r in results if r["area"] == a and r["level"] == lv]
        num = NUMBER[(a, lv)]
        md += [f"### {num} {LEVEL_NAME[lv]}: {a}", "", f"| Slno | {FIRST_COL[lv]} | Test case | Expected result | Status | Work Done |", "|---|---|---|---|---|---|"]
        for i, r in enumerate(rs, 1):
            cells = [f"{i}.", r["module"], r["test"], r["expected"], status_text(r), work_done(r)]
            md.append("| " + " | ".join(c.replace("|", "/") for c in cells) + " |")
        md.append("")
        pages = [rs[i:i + ROWS_PER_SLIDE] for i in range(0, len(rs), ROWS_PER_SLIDE)]
        for pi, chunk in enumerate(pages):
            suffix = f" ({pi + 1}/{len(pages)})" if len(pages) > 1 else ""
            table_slide(f"{num} {LEVEL_NAME[lv]}: {a}{suffix}", FIRST_COL[lv], chunk, pi * ROWS_PER_SLIDE + 1)

prs.save(out / "TEST_CASES.pptx")
(out / "TEST_CASES.md").write_text("\n".join(md), encoding="utf-8")
print("slides:", len(prs.slides), "| cases:", grand_n, "| passed:", grand_p)
